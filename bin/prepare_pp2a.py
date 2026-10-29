#!/usr/bin/env python3
"""Normalize the three supplied PP2A workbooks without inventing CAAS coordinates.

Outputs retain original rows, study/condition, reported protein-site candidates,
localization percentages and unresolved evidence. No new significance or localization
cutoff is imposed. Candidate rows sharing a record are NOT independent observations.
Requires openpyxl; original workbooks are read only.
"""
import argparse
import csv
import gzip
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

ACCESSION = r'(?:[OPQ][0-9][A-Z0-9]{3}[0-9]|[A-NR-Z][0-9](?:[A-Z][A-Z0-9]{2}[0-9]){1,2})(?:-\d+)?'
SITE = re.compile(r'([STY])(\d+)(?:\(([\d.]+)\))?')
SOURCES = [
    ('Kauko PP2A_phosphoproteome.xlsx', 'Kauko', '32071079', 'HeLa', [
        ('Dephosphorylated by PP2A act. ', 2, 'activation_decreased', 'hits'),
        ('Hyperphosphorylated by siPP2A-A', 2, 'PPP2R1A_depletion_increased', 'hits')]),
    ('Stewart PP2A-A inactivating mutation.xlsx', 'Stewart', '39939726', 'TOV21G', [
        ('Hyperphosphorylated by PP2A mut', 3, 'PPP2R1A_R183P_increased', 'hits')]),
    ('Sapkota PP2A-C KO hyperphosphorylated.xlsx', 'Sapkota', '38450154', 'HEK-293', [
        ('Hyper-phosphorylated', 1, 'PPP2CA_degradation_increased', 'hits'),
        ('Entire Data', 1, 'PPP2CA_degradation_measured', 'measured_background')]),
]


def parse_modifications(text):
    """Parse candidates by accession; retain ambiguous groups and incomplete localization."""
    matches = list(re.finditer(ACCESSION + r'(?=\s+\d+xPhospho)', text))
    candidates, unresolved = [], []
    for i, match in enumerate(matches):
        segment = text[match.end():matches[i+1].start() if i+1 < len(matches) else len(text)]
        groups = re.findall(r'(\d+)xPhospho\s*\[([^]]*)\]', segment)
        for count, content in groups:
            sites = list(SITE.finditer(content))
            # More candidates than modifications, or several mappings of one peptide,
            # must not be counted as independent, unambiguous site observations.
            ambiguous = len(matches) > 1 or len(groups) > 1 or len(sites) != int(count)
            for site in sites:
                candidates.append((match.group(), site[1], int(site[2]), site[3] or '',
                                   'ambiguous_mapping' if ambiguous else 'reported_site'))
            residue_tokens = re.findall(r'\b[STY](?:\d+)?\b', content)
            if not sites or len(residue_tokens) > len(sites):
                unresolved.append(content)
    if not candidates and not unresolved:
        unresolved.append(text or 'missing modification annotation')
    return candidates, unresolved


def prepare(source_dir, output_dir):
    import openpyxl
    output_dir.mkdir(parents=True, exist_ok=True)
    counts, inputs = Counter(), []
    fields = ['record_id','study','pmid','cell_line','sheet','source_row','collection_role','condition',
              'accession','amino_acid','protein_position','localization_percent','site_status','coordinate_space']
    with gzip.open(output_dir/'source_records.tsv.gz','wt',newline='') as raw, \
         open(output_dir/'site_candidates.tsv','w',newline='') as sites, \
         open(output_dir/'unresolved.tsv','w',newline='') as unresolved:
        rw = csv.writer(raw, delimiter='\t');rw.writerow(['record_id','source_row_json'])
        sw = csv.DictWriter(sites, fieldnames=fields, delimiter='\t');sw.writeheader()
        uw = csv.writer(unresolved, delimiter='\t');uw.writerow(['record_id','reason','original_annotation'])
        for filename, study, pmid, cell, sheets in SOURCES:
            path = source_dir/filename
            inputs.append({'file': filename, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
            # Kauko condition labels are vertically merged cells; inherit only within
            # those actual merges, rather than forward-filling arbitrary blank rows.
            book = openpyxl.load_workbook(path, read_only=study != 'Kauko', data_only=True)
            for sheet_name, header_row, condition, role in sheets:
                sheet = book[sheet_name]
                if hasattr(sheet, 'reset_dimensions'):sheet.reset_dimensions()
                rows = iter(sheet.iter_rows(values_only=True))
                for _ in range(header_row):header = next(rows)
                headers = [str(x).strip() if x is not None else f'blank_{i}' for i,x in enumerate(header)]
                counts[f'{study}/{role}/sheets'] += 1
                for row_number, values in enumerate(rows,header_row+1):
                    if not any(x is not None for x in values):continue
                    record_id = f'{study}:{sheet_name.strip()}:{row_number}'
                    # A list preserves repeated abundance headers, which a dict would lose.
                    rw.writerow([record_id,json.dumps(list(zip(headers,values)), ensure_ascii=False,default=str)])
                    d=dict(zip(headers, values));counts[f'{study}/{role}/records']+=1
                    common = dict(record_id=record_id,study=study,pmid=pmid,cell_line=cell,sheet=sheet_name,
                                  source_row=row_number,collection_role=role,condition=condition,
                                  coordinate_space='UniProt_protein_1based;CAAS_mapping_pending')
                    if study == 'Kauko':
                        modulation=d.get('Modulation')
                        if 'Modulation' in headers:
                            col=headers.index('Modulation')+1
                            for merged in sheet.merged_cells.ranges:
                                if merged.min_col <= col <= merged.max_col and merged.min_row <= row_number <= merged.max_row:
                                    modulation=sheet.cell(merged.min_row,col).value
                                    break
                            common['condition'] += ':'+str(modulation or 'unspecified')
                        accessions=re.findall(ACCESSION,str(d.get('Protein') or ''))
                        positions=[d.get(key) for key in ('1. site','2. site','3. site','4.site')]
                        candidates=[(acc,'',int(pos),'','reported_site' if len(accessions)==1 else 'ambiguous_mapping')
                                    for acc in accessions for pos in positions if isinstance(pos,(int,float))]
                        problems=[] if candidates else ['missing accession or numeric protein position']
                        annotation=str(d.get('Peptide') or '')
                    else:
                        annotation=str(d.get('Modifications in Master Proteins') or '')
                        candidates,problems=parse_modifications(annotation)
                    for acc, aa, pos, prob, status in candidates:
                        sw.writerow(dict(common,accession=acc,amino_acid=aa,protein_position=pos,
                                         localization_percent=prob,site_status=status))
                        counts[f'{study}/{role}/candidates']+=1
                        counts[f'{study}/{role}/{status}']+=1
                    for problem in problems:
                        uw.writerow([record_id,problem,annotation]);counts[f'{study}/{role}/unresolved']+=1
            book.close()
    manifest = {'inputs':inputs,'counts':dict(sorted(counts.items())),
        'interpretation':'Perturbation-associated phosphosites; not proof of direct PP2A substrates.',
        'thresholds':'Source sheets retained as supplied; no new p-value/FDR/localization cutoff.',
        'mapping':'Protein coordinates only. Validate accession/isoform and peptide against the actual reference before projecting to CAAS.',
        'conditions':'Kauko modulation labels are retained literally or inherited only inside actual merged ranges. Blank, unmerged labels stay unspecified; inspect the original blocks before assigning them.',
        'background':'Only Sapkota supplies Entire Data here; it is a measured peptide background, not all S/T residues.',
        'independence':'One source record may have several phosphosites or alternative mappings. Do not sum candidate rows as independent observations.'}
    (output_dir/'manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
    print(json.dumps(manifest['counts'],indent=2))


if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    args=parser.parse_args()
    prepare(args.source_dir,args.output_dir)
